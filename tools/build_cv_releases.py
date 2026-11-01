#!/usr/bin/env python3
"""Build RU/EN/UA Senior Unity CV previews while preserving the approved DOCX layout.

This script operates only on Git working-tree files; production release is a separate step.
Expected source: 116 nonempty paragraphs in each approved RU/EN Senior Unity DOCX.
"""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re
import subprocess
import shutil
import sys
import tempfile
from xml.sax.saxutils import escape as xml_escape

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
PUBLIC_SITE = "https://yk-developer.pages.dev/"
GAME_CASE = PUBLIC_SITE + "projects/wordbattle/"

DOC_NAMES = {
    "RU": "Yuriy_Kobylyatskiy_Senior_Unity_Developer_RU",
    "EN": "Yuriy_Kobylyatskiy_Senior_Unity_Developer_EN",
    "UA": "Yuriy_Kobylyatskiy_Senior_Unity_Developer_UA",
}

# Don't silently overwrite already approved source files.
SOURCE = {
    "RU": ROOT / "Templates/ResumeBase/RU.docx",
    "EN": ROOT / "Templates/ResumeBase/EN.docx",
}
DEST = {lang: ROOT / f"Releases/{lang}" / (name + ".docx") for lang, name in DOC_NAMES.items()}

def nonempty(doc):
    return [p for p in doc.paragraphs if p.text.strip()]

def base_style_from(p):
    first = next((r for r in p.runs if r.text.strip()), None)
    return deepcopy(first._r.rPr) if first is not None and first._r.rPr is not None else None

def write_text(p, value):
    # Preserve the original paragraph layout. Use this only for paragraphs whose
    # entire text has the same emphasis; competency items need write_competency.
    rpr = base_style_from(p)
    for node in list(p._p):
        if node.tag != qn("w:pPr"):
            p._p.remove(node)
    r = p.add_run(value)
    if rpr is not None:
        r._r.insert(0, rpr)
    return r


def write_competency(p, value):
    """Bold competency labels only, never the explanatory body text."""
    rpr = base_style_from(p)
    for node in list(p._p):
        if node.tag != qn("w:pPr"):
            p._p.remove(node)

    def add(text, bold):
        run = p.add_run(text)
        if rpr is not None:
            run._r.insert(0, deepcopy(rpr))
        run.bold = bold

    first, sep, remainder = value.partition(":")
    if not sep:
        raise RuntimeError(f"Competency item has no category label: {value[:50]}")
    add(first + ":", True)
    # The combined architecture + leadership line contains two distinct labels.
    before, marker, after = remainder.partition(" Leadership:")
    add(before, False)
    if marker:
        add(" Leadership:", True)
        add(after, False)

def clickable(p, text, url):
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    props.append(style)
    r.append(props)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)

def insert_clone_before(target, example, content, link=None):
    new_xml = deepcopy(example._p)
    target._p.addprevious(new_xml)
    p = Paragraph(new_xml, target._parent)
    write_text(p, content)
    if link:
        p.add_run(" | ")
        clickable(p, link[0], link[1])
    return p

def clean_line(s):
    s = s.strip()
    s = re.sub(r"^#{1,3}\s+", "", s)
    s = re.sub(r"^-\s+", "", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    s = s.replace("**", "").replace("\`", "")
    return s

def ua_lines():
    path = ROOT / "Drafts/UA" / (DOC_NAMES["UA"] + ".md")
    source = path.read_text(encoding="utf-8")
    active, output = True, []
    for raw in source.splitlines():
        s = raw.strip()
        if s.startswith("## Власні проєкти"):
            active = False
        if s.startswith("## Технології"):
            active = True
        if not active or not s or s == "---":
            continue
        output.append(clean_line(s))
    # Original CV has one Architecture/Leadership competency. UA editorial
    # source separates the two for readability. Merge for style alignment.
    architecture = next(i for i, t in enumerate(output) if t.startswith("Architecture & Engineering Practices:"))
    if not output[architecture + 1].startswith("Leadership:"):
        raise RuntimeError("UA architecture/leadership text unexpectedly changed")
    output[architecture] = output[architecture] + " " + output.pop(architecture + 1)
    if len(output) != 116:
        raise RuntimeError(f"UA paragraphs: expected 116, found {len(output)}; no files written")
    for t in ("з 2007 року", "з 2011 року", "SOLID", "Zenject", "UniTask"):
        if not any(t in x for x in output):
            raise RuntimeError(f"UA marker {t!r} missing")
    return output

def change_core(doc, lang):
    ps = nonempty(doc)
    if len(ps) != 116:
        raise RuntimeError(f"{lang}: expected 116 paragraphs, got {len(ps)}")
    if lang == "UA":
        translated = ua_lines()
        for idx, (p, value) in enumerate(zip(ps, translated, strict=True)):
            if 28 <= idx <= 34:
                write_competency(p, value)
            else:
                write_text(p, value)
        # Make website and contact links actually clickable (not just visible text).
        write_text(ps[3], "Україна | Лише віддалена робота | +380957575677 (Telegram, Viber) | finmagnat@gmail.com | ")
        clickable(ps[3], "Портфоліо", PUBLIC_SITE)
        write_text(ps[4], "GitHub: ")
        clickable(ps[4], "github.com/finmagnat", "https://github.com/finmagnat")
        ps[4].add_run(" | LinkedIn: ")
        clickable(ps[4], "yuriy-kobylyatskiy", "https://www.linkedin.com/in/yuriy-kobylyatskiy")
    else:
        summary = ("В разработке ПО с 2007 года, в GameDev — с 2011 года." if lang == "RU" else
                   "Software development since 2007; GameDev since 2011.")
        write_text(ps[6], summary)
        write_competency(ps[28], "Unity Development: Unity 2D/3D, gameplay programming, UI programming, animation systems, Unity Addressables, UniTask, async/await, CancellationToken, ScriptableObject, DOTween, Spine, Steamworks, custom editor tools.")
        write_competency(ps[32], "Optimization & Automation: Memory optimization, build size optimization, performance optimization, crash reduction, Unity Memory Profiler, CI/CD (Jenkins, TeamCity), Bash, Rclone, Cron.")
        practices = {
            "RU": "Architecture & Engineering Practices: Dependency Injection (Zenject), SOLID, KISS, DRY, YAGNI, модульная архитектура, рефакторинг, техническое проектирование. Leadership: оценка задач, наставничество, управление командами, найм, техническая документация, аудит проектов.",
            "EN": "Architecture & Engineering Practices: Dependency Injection (Zenject), SOLID, KISS, DRY, YAGNI, modular architecture, refactoring, technical design. Leadership: estimation, mentoring, team leadership, hiring, technical documentation, project audits.",
        }
        write_competency(ps[34], practices[lang])
    return ps

def add_wordbattle(doc, lang):
    ps = nonempty(doc)
    if len(ps) != 116:
        raise RuntimeError(f"{lang}: unsupported paragraph count after edits")
    # Heading, job title, metadata, bullet and technology paragraph exemplars.
    target = ps[102]
    if lang == "RU":
        content = [
            ("Собственные проекты (Independent Projects)", ps[102], None),
            ("WordBattle — собственный инди-проект", ps[36], None),
            ("Unity / Android | В разработке и тестировании", ps[37], ("Кейс в портфолио", GAME_CASE)),
            ("Разрабатываю словесную стратегическую игру с AI-соперником. Отвечаю за архитектуру, алгоритмы поиска слов, игровые механики, UI, интеграцию SDK и подготовку Android-релиза.", ps[38], None),
            ("Реализую игровые сервисы, аналитику и монетизацию с использованием PlayFab, GameAnalytics, Google Play Billing и AdMob.", ps[39], None),
            ("Технологии: Unity, C#, Unity Addressables, Zenject, UniTask, PlayFab, GameAnalytics, Google Play Billing, AdMob.", ps[44], None),
        ]
    elif lang == "EN":
        content = [
            ("Independent Projects", ps[102], None),
            ("WordBattle — independent indie game", ps[36], None),
            ("Unity / Android | In development and testing", ps[37], ("Portfolio case study", GAME_CASE)),
            ("Building a word strategy game against an AI opponent, owning architecture, word-search algorithms, gameplay, UI, SDK integrations and Android release preparation.", ps[38], None),
            ("Implementing game services, analytics and monetization with PlayFab, GameAnalytics, Google Play Billing and AdMob.", ps[39], None),
            ("Technologies: Unity, C#, Unity Addressables, Zenject, UniTask, PlayFab, GameAnalytics, Google Play Billing, AdMob.", ps[44], None),
        ]
    else:
        content = [
            ("Власні проєкти (Independent Projects)", ps[102], None),
            ("WordBattle — власний інді-проєкт", ps[36], None),
            ("Unity / Android | У розробці та тестуванні", ps[37], ("Опис проєкту в портфоліо", GAME_CASE)),
            ("Розробляю словесну стратегічну гру з AI-суперником, відповідаю за архітектуру, алгоритми пошуку слів, ігрові механіки, UI, інтеграцію SDK та підготовку Android-релізу.", ps[38], None),
            ("Реалізую ігрові сервіси, аналітику та монетизацію з використанням PlayFab, GameAnalytics, Google Play Billing і AdMob.", ps[39], None),
            ("Технології: Unity, C#, Unity Addressables, Zenject, UniTask, PlayFab, GameAnalytics, Google Play Billing, AdMob.", ps[44], None),
        ]
    for value, exemplar, link in content:
        insert_clone_before(target, exemplar, value, link=link)

def render_pdf(docx, pdf_dir):
    if not shutil.which("libreoffice"):
        raise RuntimeError("libreoffice missing: cannot export PDFs")
    proc = subprocess.run(["libreoffice", "-env:UserInstallation=file:///tmp/resume-lo-profile",
                           "--headless", "--convert-to", "pdf:writer_pdf_Export",
                           "--outdir", str(pdf_dir), str(docx)],
                          check=True, text=True, capture_output=True, timeout=150)
    pdf = pdf_dir / docx.with_suffix(".pdf").name
    if not pdf.is_file() or pdf.stat().st_size < 10000:
        raise RuntimeError(f"PDF export missing: {pdf}; stdout={proc.stdout} stderr={proc.stderr}")
    return pdf

def validate(docx, pdf, lang):
    from pypdf import PdfReader
    d = Document(docx)
    txt = "\n".join(p.text for p in d.paragraphs)
    for token in ("WordBattle", "2007", "2011", "Zenject", "Addressables", "UniTask", "SOLID", "KISS", "DRY", "YAGNI", "Jenkins", "TeamCity"):
        if token not in txt:
            raise RuntimeError(f"{lang}: DOCX missing {token}")
    if "VContainer" in txt or "github.com/finmagnat/WordBattle" in txt:
        raise RuntimeError(f"{lang}: unexpected technology claim or private game repository URL")
    urls = [r.target_ref for r in d.part.rels.values() if r.reltype == RT.HYPERLINK]
    if GAME_CASE not in urls or PUBLIC_SITE not in urls:
        raise RuntimeError(f"{lang}: missing site or WordBattle actual hyperlink relationship: {urls}")
    reader = PdfReader(str(pdf))
    if not (2 <= len(reader.pages) <= 9):
        raise RuntimeError(f"{lang}: suspicious page count = {len(reader.pages)}")
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    for token in ("WordBattle", "2007", "2011", "SOLID"):
        if token not in pdf_text:
            raise RuntimeError(f"{lang}: PDF missing {token}")
    annotations = []
    for page in reader.pages:
        for a in page.get("/Annots") or []:
            annotation = a.get_object()
            action = annotation.get("/A")
            if action and action.get("/URI"):
                annotations.append(str(action["/URI"]))
    if GAME_CASE not in annotations or PUBLIC_SITE not in annotations:
        raise RuntimeError(f"{lang}: PDF URL annotation missing: {annotations}")
    print(f"{lang}: DOCX={docx.stat().st_size} bytes, PDF={pdf.stat().st_size} bytes, pages={len(reader.pages)}, links OK")

def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-releases", action="store_true",
                        help="Explicitly replace the approved Releases files; review before running")
    args = parser.parse_args()
    # Never overwrite the owner's manually polished releases by default.
    outputs = DEST if args.output_releases else {
        lang: ROOT / ".artifacts/cv-preview" / lang / (name + ".docx")
        for lang, name in DOC_NAMES.items()
    }
    ua_base = ua_lines()
    made = []
    # RU/EN sources are read before writing any output so the two languages stay consistent.
    documents = {"RU": Document(SOURCE["RU"]), "EN": Document(SOURCE["EN"]), "UA": Document(SOURCE["RU"])}
    for lang in ("RU", "EN", "UA"):
        doc = documents[lang]
        change_core(doc, lang)
        add_wordbattle(doc, lang)
        output = outputs[lang]
        output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output)
        made.append(output)
    for lang, output in zip(("RU", "EN", "UA"), made, strict=True):
        pdf = render_pdf(output, output.parent)
        validate(output, pdf, lang)
    print("All three CV DOCX/PDF previews generated; owner should visually review exported pages before publishing.")

if __name__ == "__main__":
    main()
